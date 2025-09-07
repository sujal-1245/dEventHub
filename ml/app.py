from fastapi import FastAPI, UploadFile, File
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
import shutil, os, json
from dotenv import load_dotenv
import requests
from io import BytesIO
from fpdf import FPDF
from docx import Document
import fitz  # PyMuPDF for PDF hyperlink extraction

# ---------------- Load env ----------------
load_dotenv()
PORT = int(os.getenv("PORT", 8000))
UPLOAD_FOLDER = os.getenv("UPLOAD_FOLDER", "uploads")
RESUME_TEMPLATES_FOLDER = os.getenv("RESUME_TEMPLATES_FOLDER", "templates")
os.makedirs(UPLOAD_FOLDER, exist_ok=True)
os.makedirs(RESUME_TEMPLATES_FOLDER, exist_ok=True)

GROQ_API_KEY = os.getenv("GROQ_API_KEY")
GROQ_API_URL = "https://api.groq.com/openai/v1/chat/completions"

# ---------------- FastAPI app ----------------
app = FastAPI(title="dEventHub AI/ML Service")

# ---------------- CORS ----------------
origins = [
    "http://localhost:5173",
    "http://127.0.0.1:5173",
    "https://eventhub-git-main-sujal-1245s-projects.vercel.app"
]

app.add_middleware(
    CORSMiddleware,
    allow_origins=origins,
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


# ---------------- Models ----------------
class ResumeTemplateRequest(BaseModel):
    template_name: str

class ChatbotRequest(BaseModel):
    message: str

class RecommendRequest(BaseModel):
    query: str = ""      # optional search filter
    status: str = "open" # can be "open" or "upcoming"

# ---------------- Helper functions ----------------
def extract_text_from_docx(file_path: str):
    doc = Document(file_path)
    text = ""
    links = {}
    for para in doc.paragraphs:
        text += para.text + "\n"
        for run in para.runs:
            if hasattr(run, "hyperlink") and run.hyperlink:
                url = run.hyperlink.target
                text_fragment = run.text.strip()
                if "linkedin" in url.lower() or "linkedin" in text_fragment.lower():
                    links["linkedin"] = url
                elif "github" in url.lower() or "github" in text_fragment.lower():
                    links["github"] = url
                else:
                    if "portfolio" not in links:
                        links["portfolio"] = url
    return text, links

def extract_text_from_pdf(file_path: str):
    text = ""
    links = {}
    doc = fitz.open(file_path)
    for page in doc:
        text += page.get_text() + "\n"
        for link in page.get_links():
            if "uri" in link:
                url = link["uri"]
                if "linkedin" in url.lower() and "linkedin" not in links:
                    links["linkedin"] = url
                elif "github" in url.lower() and "github" not in links:
                    links["github"] = url
                elif "portfolio" not in links:
                    links["portfolio"] = url
    return text, links

def extract_text_from_file(file_path: str):
    """Extract text and links from DOCX or PDF resume"""
    text, links = "", {}
    if file_path.lower().endswith(".docx"):
        text, links = extract_text_from_docx(file_path)
    elif file_path.lower().endswith(".pdf"):
        text, links = extract_text_from_pdf(file_path)
    else:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            text = f.read()
    print(f"✅ Extracted text length: {len(text)} chars, links found: {links}")
    return text, links

def analyze_resume_with_ai(text: str, links: dict = {}) -> dict:
    if not GROQ_API_KEY:
        return {"error": "GROQ_API_KEY not configured."}

    prompt = f"""
You are an AI resume analyzer. Read the following resume text and any hyperlinks provided, then return a JSON with these fields:
- name
- contact (phone)
- email
- linkedin
- github
- portfolio
- skills (list)
- education (list of dicts: degree, university, duration, location, GPA if available)
- experience (list of dicts: title, company, duration, description, achievements)
- score (0-100)
- feedback (list of suggestions to improve ATS compatibility)

Resume Text:
{text}

Hyperlinks extracted from resume:
{json.dumps(links)}

Only return valid JSON. Do not include extra text.
"""

    headers = {"Authorization": f"Bearer {GROQ_API_KEY}", "Content-Type": "application/json"}
    payload = {
        "model": "meta-llama/llama-4-maverick-17b-128e-instruct",
        "messages": [{"role": "user", "content": prompt}],
        "max_tokens": 1500
    }

    try:
        response = requests.post(GROQ_API_URL, headers=headers, json=payload)
        response.raise_for_status()
        data = response.json()
        content = data.get("choices", [{}])[0].get("message", {}).get("content", "{}")
        first_brace = content.find("{")
        last_brace = content.rfind("}")
        json_str = content[first_brace:last_brace+1] if first_brace != -1 and last_brace != -1 else "{}"
        result = json.loads(json_str)
        print("✅ AI analysis result keys:", result.keys())
        return result
    except Exception as e:
        print("❌ AI analysis error:", e)
        return {"error": str(e)}

def build_ats_resume(extracted_info: dict, template_name: str) -> bytes:
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", 'B', 16)
    pdf.cell(0, 10, extracted_info.get("name", "Your Name"), ln=True)
    pdf.set_font("Arial", '', 12)

    for field in ["contact", "email", "linkedin", "github", "portfolio"]:
        if extracted_info.get(field):
            pdf.cell(0, 10, f"{field.capitalize()}: {extracted_info[field]}", ln=True)

    pdf.ln(5)
    pdf.set_font("Arial", 'B', 14)
    pdf.cell(0, 10, "Skills", ln=True)
    pdf.set_font("Arial", '', 12)
    pdf.multi_cell(0, 8, ", ".join(extracted_info.get("skills", [])))

    pdf.ln(3)
    pdf.set_font("Arial", 'B', 14)
    pdf.cell(0, 10, "Education", ln=True)
    pdf.set_font("Arial", '', 12)
    for edu in extracted_info.get("education", []):
        edu_str = f"{edu.get('degree','')} @ {edu.get('university','')} ({edu.get('duration','')})"
        if edu.get("GPA"): edu_str += f" - GPA: {edu.get('GPA')}"
        pdf.multi_cell(0, 8, edu_str)

    pdf.ln(3)
    pdf.set_font("Arial", 'B', 14)
    pdf.cell(0, 10, "Experience", ln=True)
    pdf.set_font("Arial", '', 12)
    for exp in extracted_info.get("experience", []):
        exp_str = f"{exp.get('title','')} @ {exp.get('company','')} ({exp.get('duration','')})"
        pdf.multi_cell(0, 8, exp_str)
        if exp.get("description"):
            pdf.multi_cell(0, 8, f"  {exp['description']}")
        for ach in exp.get("achievements", []):
            pdf.multi_cell(0, 8, f"   - {ach}")

    buffer = BytesIO()
    pdf.output(buffer)
    buffer.seek(0)
    return buffer.read()

# ---------------- Routes ----------------
@app.post("/resume")
def resume_analysis(file: UploadFile = File(...)):
    file_path = f"{UPLOAD_FOLDER}/{file.filename}"
    with open(file_path, "wb") as buffer:
        shutil.copyfileobj(file.file, buffer)
    print(f"📄 File saved: {file_path}")

    text, links = extract_text_from_file(file_path)
    result = analyze_resume_with_ai(text, links)
    print("✅ Returning AI resume analysis response")
    return {"analysis": result}

@app.post("/resume/build")
def build_resume(req: ResumeTemplateRequest, file: UploadFile = File(...)):
    file_path = f"{UPLOAD_FOLDER}/{file.filename}"
    with open(file_path, "wb") as buffer:
        shutil.copyfileobj(file.file, buffer)
    print(f"📄 File saved for building ATS resume: {file_path}")

    text, links = extract_text_from_file(file_path)
    extracted_info = analyze_resume_with_ai(text, links)
    pdf_bytes = build_ats_resume(extracted_info, req.template_name)

    print(f"✅ ATS resume PDF built, size: {len(pdf_bytes)} bytes")
    return {
        "filename": f"ats_resume_{file.filename}.pdf",
        "content": pdf_bytes.hex()
    }

@app.post("/chat")
def chatbot(req: ChatbotRequest):
    if not GROQ_API_KEY:
        return {"response": "API key not set. Please configure GROQ_API_KEY in .env."}
    headers = {"Authorization": f"Bearer {GROQ_API_KEY}", "Content-Type": "application/json"}
    payload = {
        "model": "meta-llama/llama-4-maverick-17b-128e-instruct",
        "messages": [{"role": "user", "content": req.message}],
        "max_tokens": 150
    }
    try:
        response = requests.post(GROQ_API_URL, headers=headers, json=payload)
        response.raise_for_status()
        data = response.json()
        reply = data.get("choices", [{}])[0].get("message", {}).get("content", "Sorry, no response.")
        return {"response": reply}
    except Exception as e:
        return {"response": f"Error: {str(e)}"}

@app.post("/recommend")
def recommend_events(req: RecommendRequest):
    """
    Fetch live hackathons/events from Devpost API.
    """
    try:
        url = f"https://devpost.com/api/hackathons?status={req.status}&page=1"
        response = requests.get(url, timeout=10)
        response.raise_for_status()
        data = response.json()

        events = []
        for h in data.get("hackathons", []):
            event = {
                "title": h.get("title"),
                "location": h.get("location", "Online"),
                "url": h.get("url"),
                "submission_deadline": h.get("submission_period_dates"),
                "thumbnail": h.get("thumbnail_url")
            }
            if req.query:
                if req.query.lower() in (event["title"] or "").lower():
                    events.append(event)
            else:
                events.append(event)

        return {"events": events}

    except Exception as e:
        return {"error": str(e)}

@app.get("/")
def root():
    return {"message": "dEventHub ML Service running"}
